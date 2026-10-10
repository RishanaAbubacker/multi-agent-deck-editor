from pptx import Presentation
from document.pptx_inspector import inspect_pptx


def test_inspector_extracts_slide_and_text(tmp_path):
    path = tmp_path / "sample.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    shape = slide.shapes.add_textbox(100, 100, 3000000, 500000)
    shape.text = "Hello"
    prs.save(path)
    deck = inspect_pptx(path)
    assert len(deck.slides) == 1
    assert any(s.text == "Hello" for s in deck.slides[0].shapes)
