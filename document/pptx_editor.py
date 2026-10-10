from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from schemas.models import EditAction, OperationResult

SUPPORTED_OPERATIONS = {
    "replace_text": "parameters: text (string)",
    "set_font_size": "parameters: font_size_pt (positive number)",
    "set_bold": "parameters: value (boolean)",
    "set_italic": "parameters: value (boolean)",
    "set_font_name": "parameters: font_name (string)",
    "set_font_color": "parameters: hex_color (RRGGBB string)",
    "set_shape_position": "parameters: left, top (integer EMU units)",
    "set_shape_size": "parameters: width, height (positive integer EMU units)",
}


def _find_shape(prs, slide_number: int, shape_id: int):
    if slide_number < 1 or slide_number > len(prs.slides):
        raise ValueError(f"Slide number out of range: {slide_number}")
    slide = prs.slides[slide_number - 1]
    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape
    raise ValueError(f"Shape ID {shape_id} not found on slide {slide_number}")


def _set_all_runs(shape, attribute: str, value):
    if not getattr(shape, "has_text_frame", False):
        raise ValueError("Shape does not contain a text frame")
    found = False
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            setattr(run.font, attribute, value)
            found = True
    if not found:
        raise ValueError("Shape has no text runs to format")


def apply_edit_plan(input_path: Path, output_path: Path, actions: list[EditAction]) -> list[OperationResult]:
    prs = Presentation(str(input_path))
    results = []
    for action in actions:
        try:
            shape = _find_shape(prs, action.slide_number, action.shape_id)
            p = action.parameters
            if action.operation == "replace_text":
                if not getattr(shape, "has_text_frame", False):
                    raise ValueError("Shape does not contain text")
                # Replacing all text can affect run formatting; use only when the plan requires a full text replacement.
                shape.text_frame.text = str(p["text"])
            elif action.operation == "set_font_size":
                size = float(p["font_size_pt"])
                if size <= 0:
                    raise ValueError("font_size_pt must be positive")
                _set_all_runs(shape, "size", __import__("pptx").util.Pt(size))
            elif action.operation == "set_bold":
                _set_all_runs(shape, "bold", bool(p["value"]))
            elif action.operation == "set_italic":
                _set_all_runs(shape, "italic", bool(p["value"]))
            elif action.operation == "set_font_name":
                _set_all_runs(shape, "name", str(p["font_name"]))
            elif action.operation == "set_font_color":
                raw = str(p["hex_color"]).lstrip("#")
                if len(raw) != 6:
                    raise ValueError("hex_color must be six hexadecimal characters")
                color = RGBColor.from_string(raw)
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = color
            elif action.operation == "set_shape_position":
                shape.left, shape.top = int(p["left"]), int(p["top"])
            elif action.operation == "set_shape_size":
                width, height = int(p["width"]), int(p["height"])
                if width <= 0 or height <= 0:
                    raise ValueError("width and height must be positive")
                shape.width, shape.height = width, height
            else:
                raise ValueError(f"Unsupported operation: {action.operation}")
            results.append(OperationResult(action_id=action.action_id, issue_id=action.issue_id, success=True,
                                           message="Operation applied", slide_number=action.slide_number, shape_id=action.shape_id))
        except Exception as exc:
            results.append(OperationResult(action_id=action.action_id, issue_id=action.issue_id, success=False,
                                           message=str(exc), slide_number=action.slide_number, shape_id=action.shape_id))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return results
