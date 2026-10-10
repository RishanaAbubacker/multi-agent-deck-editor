from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from schemas.models import DeckInfo, SlideInfo, ShapeInfo


def _color_value(font) -> str | None:
    try:
        color = font.color
        if color and color.type is not None and color.rgb is not None:
            return str(color.rgb)
    except (AttributeError, TypeError, ValueError):
        return None
    return None


def inspect_pptx(path: Path) -> DeckInfo:
    prs = Presentation(str(path))
    slides = []
    for slide_no, slide in enumerate(prs.slides, start=1):
        shapes = []
        for shape in slide.shapes:
            runs = []
            if getattr(shape, "has_text_frame", False):
                for para_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                    for run_idx, run in enumerate(paragraph.runs):
                        runs.append({
                            "paragraph_index": para_idx,
                            "run_index": run_idx,
                            "text": run.text,
                            "font_name": run.font.name,
                            "font_size_pt": run.font.size.pt if run.font.size else None,
                            "bold": run.font.bold,
                            "italic": run.font.italic,
                            "font_color": _color_value(run.font),
                            "alignment": str(paragraph.alignment) if paragraph.alignment is not None else None,
                        })
            shapes.append(ShapeInfo(
                shape_id=shape.shape_id,
                name=shape.name,
                shape_type=str(shape.shape_type),
                left=shape.left,
                top=shape.top,
                width=shape.width,
                height=shape.height,
                text=shape.text if getattr(shape, "has_text_frame", False) else "",
                font_runs=runs,
            ))
        slides.append(SlideInfo(slide_number=slide_no, layout_name=slide.slide_layout.name, shapes=shapes))
    return DeckInfo(path=str(path), slide_width=prs.slide_width, slide_height=prs.slide_height, slides=slides)
