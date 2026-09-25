from pathlib import Path
from pptx import Presentation

class EditorAgent:
    name = "editor_agent"

    def run(self, input_path: Path, output_path: Path, plan):
        prs = Presentation(str(input_path))
        changes = []
        for action in plan.actions:
            if action.action != "strip_text_whitespace":
                continue
            slide = prs.slides[action.slide_number - 1]
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                old = shape.text
                new = old.strip()
                if old != new:
                    shape.text_frame.clear()
                    shape.text_frame.paragraphs[0].text = new
                    changes.append({
                        "action_id": action.action_id,
                        "slide_number": action.slide_number,
                        "shape_id": shape.shape_id,
                        "action": action.action,
                        "old_value": old,
                        "new_value": new
                    })
        prs.save(str(output_path))
        return changes
